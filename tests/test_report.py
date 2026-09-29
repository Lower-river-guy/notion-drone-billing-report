from datetime import date, datetime
from pathlib import Path


def _as_date(value):
    if isinstance(value, datetime):
        return value.date()
    return value

from openpyxl import load_workbook

from drone_billing.billing import BillableFlight, ReviewItem
from drone_billing.cost_config import cost_for_acres
from drone_billing.report import (
    BILLING_REPORT_HEADERS,
    FLIGHT_DETAIL_HEADERS,
    HEADER_ROW,
    TOTAL_LABEL,
    billing_title,
    build_workbook,
    report_filename,
    save_workbook,
)


def _billable(**kwargs) -> BillableFlight:
    acres = kwargs.pop("acres", 200)
    defaults = {
        "page_id": "p1",
        "page_url": "https://app.notion.com/p1",
        "title": "1575-Mid Valley",
        "project_number": "1575",
        "project_name": "Mid Valley",
        "flight_date": date(2026, 9, 28),
        "flight_type": "Progress",
        "drone_equipment": "Wingtra Ray",
        "acres": acres,
        "cost": cost_for_acres(acres),
        "status": "Completed",
    }
    defaults.update(kwargs)
    return BillableFlight(**defaults)


def test_billing_title():
    assert billing_title("2026-09") == "DRONE BILLINGS - September 2026"


def test_workbook_generation(tmp_path: Path):
    billable = [
        _billable(
            page_id="p-late",
            project_number="1582",
            project_name="Victorville LF Phase 1B",
            title="1582-Victorville",
            flight_date=date(2026, 9, 24),
            acres=300,
            status="Completed",
        ),
        _billable(
            page_id="p-wip",
            flight_date=date(2026, 9, 28),
            status="In Process",
        ),
        _billable(
            page_id="p-early",
            project_number="1582",
            project_name="Victorville LF Phase 1B",
            title="1582-Victorville",
            flight_date=date(2026, 9, 3),
            acres=300,
            status="Completed",
        ),
        _billable(page_id="p-done", flight_date=date(2026, 9, 2), status="Completed"),
    ]
    wb = build_workbook("2026-09", billable, [])
    path = save_workbook(wb, tmp_path / report_filename("2026-09"))

    loaded = load_workbook(path)
    assert loaded.sheetnames == ["Billing Report", "Flight Detail", "Review Required"]

    ws = loaded["Billing Report"]
    assert ws.cell(row=1, column=1).value == "DRONE BILLINGS - September 2026"
    assert ws.cell(row=1, column=1).alignment.horizontal == "center"
    assert "A1:I1" in ws.merged_cells
    assert ws.cell(row=2, column=1).value is None

    headers = [ws.cell(row=HEADER_ROW, column=col).value for col in range(1, 10)]
    assert headers == BILLING_REPORT_HEADERS
    assert headers[6] == "STATUS"
    assert all("Notion" not in str(h) for h in headers)
    assert all("PAGE" not in str(h).upper() or h == "$ BILLED AMOUNT" for h in headers)
    assert "ID" not in " ".join(str(h) for h in headers)
    assert "URL" not in " ".join(str(h) for h in headers)
    for col in range(1, 10):
        assert ws.cell(row=HEADER_ROW, column=col).font.bold is True

    # Grouped by project number then date, blank row between project groups.
    assert ws.cell(row=4, column=1).value == "1575"
    assert _as_date(ws.cell(row=4, column=3).value) == date(2026, 9, 2)
    assert ws.cell(row=5, column=1).value == "1575"
    assert ws.cell(row=5, column=7).value == "In Process"
    assert float(ws.cell(row=5, column=9).value) == 2500.00
    assert ws.cell(row=6, column=1).value is None
    assert ws.cell(row=7, column=1).value == "1582"
    assert _as_date(ws.cell(row=7, column=3).value) == date(2026, 9, 3)
    assert ws.cell(row=8, column=1).value == "1582"
    assert _as_date(ws.cell(row=8, column=3).value) == date(2026, 9, 24)
    assert ws.cell(row=9, column=1).value is None
    assert ws.cell(row=10, column=1).value == TOTAL_LABEL
    assert float(ws.cell(row=10, column=9).value) == 2500.00 + 2500.00 + 2930.68 + 2930.68
    assert ws.cell(row=10, column=1).font.bold is True
    assert ws.cell(row=10, column=9).font.bold is True
    assert ws.cell(row=4, column=9).number_format == '"$"#,##0.00'
    assert ws.cell(row=4, column=3).number_format == "MM/DD/YYYY"
    assert ws.cell(row=4, column=6).number_format == "#,##0"

    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.fitToWidth == 1
    assert ws.print_title_rows in {"3:3", "$3:$3"}

    detail = loaded["Flight Detail"]
    assert [detail.cell(row=1, column=col).value for col in range(1, 12)] == FLIGHT_DETAIL_HEADERS
    assert detail.cell(row=2, column=7).value in {"Completed", "In Process"}
    assert detail.cell(row=2, column=10).value
    assert str(detail.cell(row=2, column=11).value).startswith("https://")
    assert float(detail.cell(row=2, column=9).value) in {2500.00, 2930.68}


def test_review_sheet(tmp_path: Path):
    review = [
        ReviewItem(
            page_id="r1",
            page_url="https://app.notion.com/r1",
            title="TEMPLATE x",
            reason="Title starts with TEMPLATE (not billable)",
            status="Completed",
        )
    ]
    wb = build_workbook("2026-09", [], review)
    path = save_workbook(wb, tmp_path / "test.xlsx")
    loaded = load_workbook(path)
    ws = loaded["Review Required"]
    assert ws.max_row >= 2
    assert "TEMPLATE" in str(ws.cell(row=2, column=6).value)
    assert ws.cell(row=2, column=7).value == "r1"
