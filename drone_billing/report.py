"""Excel report generation via openpyxl."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from drone_billing.billing import BillableFlight, ReviewItem, SummaryRow


def report_filename(billing_month: str) -> str:
    return f"Drone_Billing_Report_{billing_month}.xlsx"


def _money(value: Decimal | None) -> float | str:
    if value is None:
        return ""
    return float(value)


def build_workbook(
    billing_month: str,
    summary_rows: list[SummaryRow],
    billable: list[BillableFlight],
    review: list[ReviewItem],
) -> Workbook:
    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Monthly Summary"

    sum_headers = [
        "Billing Month",
        "Project Number",
        "Project Name",
        "Acres",
        "Completed Flights",
        "Cost Per Flight",
        "Total Billing Cost",
        "Review Note",
    ]
    ws_sum.append(sum_headers)
    month_total = Decimal("0")
    for row in summary_rows:
        ws_sum.append(
            [
                row.billing_month,
                row.project_number,
                row.project_name,
                row.acres,
                row.completed_flights,
                _money(row.cost_per_flight),
                _money(row.total_billing_cost),
                row.review_note or "",
            ]
        )
        month_total += row.total_billing_cost

    ws_sum.append([])
    total_row = ws_sum.max_row + 1
    ws_sum.append(
        [
            "MONTH TOTAL",
            "",
            "",
            "",
            "",
            "",
            _money(month_total),
            "",
        ]
    )
    for cell in ws_sum[total_row]:
        if cell.value:
            cell.font = Font(bold=True)

    ws_detail = wb.create_sheet("Flight Detail")
    detail_headers = [
        "Flight Date",
        "Project Number",
        "Project Name",
        "Flight Type",
        "Drone Equipment",
        "Acres",
        "Cost Tier",
        "Flight Cost",
        "Notion Page URL",
    ]
    ws_detail.append(detail_headers)
    for b in sorted(billable, key=lambda x: (x.flight_date, x.project_number)):
        assert b.cost.total is not None
        ws_detail.append(
            [
                b.flight_date.isoformat(),
                b.project_number,
                b.project_name,
                b.flight_type or "",
                b.drone_equipment or "",
                b.acres,
                b.cost.label,
                _money(b.cost.total),
                b.page_url,
            ]
        )

    ws_rev = wb.create_sheet("Review Required")
    rev_headers = [
        "Title",
        "Project Number",
        "Project Name",
        "Flight Date",
        "Reason",
        "Notion Page URL",
    ]
    ws_rev.append(rev_headers)
    for r in sorted(review, key=lambda x: x.title):
        ws_rev.append(
            [
                r.title,
                r.project_number,
                r.project_name,
                r.flight_date.isoformat() if r.flight_date else "",
                r.reason,
                r.page_url,
            ]
        )

    for ws in (ws_sum, ws_detail, ws_rev):
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18

    return wb


def save_workbook(wb: Workbook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
